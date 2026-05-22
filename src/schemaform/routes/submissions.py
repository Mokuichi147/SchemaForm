from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime, time
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    Response,
)

from schemaform.calculated import evaluate_formula
from schemaform.file_signing import file_url_builder
from schemaform.fields import (
    clean_empty_recursive,
    expand_group_array_rows,
    flatten_fields,
    flatten_filter_fields,
    format_array_group_value,
    get_nested_value,
    set_nested_value,
)
from schemaform.filters import (
    apply_filters,
    collect_file_ids,
    normalize_number,
    parse_bool,
    resolve_file_infos,
    resolve_file_names,
    value_to_text,
)
from schemaform.master import (
    build_master_reference_context,
    enrich_master_options,
    validate_master_references,
)
from schemaform.schema import fields_from_schema
from schemaform.utils import new_ulid, now_utc

router = APIRouter()


async def admin_guard(request: Request) -> None:
    await request.app.state.auth_provider.require_admin(request)


async def form_editor_guard(request: Request, form_id: str) -> None:
    """フォームの作成グループに所属するユーザー（または管理者）を許可する。"""
    from schemaform.app import can_edit_form

    await request.app.state.auth_provider.require_login(request)
    storage = request.app.state.storage
    form = storage.forms.get_form(form_id)
    if not form:
        raise HTTPException(status_code=404, detail="フォームが見つかりません")
    if not can_edit_form(request, form):
        raise HTTPException(
            status_code=403, detail="このフォームを管理する権限がありません"
        )


async def resolve_user_display_map(request: Request) -> dict[int, str]:
    """user_id → 表示名 のマップを認証プロバイダから構築する。"""
    user_display_map: dict[int, str] = {}
    auth = request.app.state.auth_provider
    current_user = getattr(request.state, "current_user", None)
    list_users = getattr(auth, "list_users", None)
    if list_users is None:
        return user_display_map
    try:
        users = await list_users((current_user or {}).get("token", ""))
    except Exception:
        users = []
    for u in users:
        uid = u.get("id")
        if uid is not None:
            user_display_map[uid] = u.get("display_name") or u.get("username") or ""
    return user_display_map


def resolve_user_label(item: dict[str, Any], user_display_map: dict[int, str]) -> str:
    uid = item.get("user_id")
    if uid in user_display_map and user_display_map[uid]:
        return user_display_map[uid]
    return item.get("username") or ""


def build_submission_display_columns(
    storage: Any, fields: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], dict[str, dict[str, dict[str, Any]]]]:
    flat_fields = flatten_fields(fields, expand_rows_for_group_arrays=True)
    display_columns: list[dict[str, Any]] = []
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]] = {}

    for field in flat_fields:
        flat_key = field["flat_key"]
        if field.get("type") != "master":
            display_columns.append(
                {
                    "kind": "default",
                    "label": field["flat_label"],
                    "field": field,
                }
            )
            continue

        context = build_master_reference_context(storage, field)
        lookup = {row["id"]: row for row in context["records"]}
        for row in context["records"]:
            rid = row["id"]
            if ":" in rid:
                base_id = rid.rsplit(":", 1)[0]
                if base_id not in lookup:
                    lookup[base_id] = row
        master_lookup_by_field[flat_key] = lookup
        display_items = context["display_items"]

        # フォーム参照の選択値そのものは常に列として表示する。
        display_columns.append(
            {
                "kind": "master_label",
                "label": field["flat_label"],
                "field": field,
            }
        )

        if display_items:
            for item in display_items:
                display_columns.append(
                    {
                        "kind": "master_display",
                        "label": f"{field['flat_label']}.{item['label']}",
                        "field": field,
                        "display_key": item["key"],
                        "display_type": item.get("type", ""),
                    }
                )

    return display_columns, master_lookup_by_field


def render_master_display_text(
    raw_value: Any,
    lookup: dict[str, dict[str, Any]],
    display_key: str | None = None,
) -> str:
    def resolve_one(value: Any) -> str:
        if value in (None, ""):
            return ""
        row = lookup.get(str(value))
        if not row:
            return ""
        if display_key:
            return str((row.get("values") or {}).get(display_key, ""))
        return str(row.get("label", ""))

    if isinstance(raw_value, list):
        parts = [text for text in (resolve_one(item) for item in raw_value) if text]
        return ", ".join(parts)
    return resolve_one(raw_value)


def _get_column_sort_key(
    item: dict[str, Any],
    column: dict[str, Any],
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]],
) -> tuple[int, Any]:
    """Return a comparable sort key for a single submission row by column."""
    data = item.get("data_json", {})
    field = column["field"]
    flat_key = field["flat_key"]
    value = get_nested_value(data, flat_key)
    is_numeric = field.get("type") in ("number", "integer")

    if value is None or value == "":
        return (1, 0.0) if is_numeric else (1, "")

    if isinstance(value, list):
        if field.get("type") == "master":
            lookup = master_lookup_by_field.get(flat_key, {})
            display_key = column.get("display_key")
            if display_key:
                text = render_master_display_text(value, lookup, str(display_key))
            else:
                text = render_master_display_text(value, lookup)
            return (0, text.lower()) if text else (1, "")
        if is_numeric:
            try:
                return (0, sum(float(v) for v in value))
            except (ValueError, TypeError):
                return (1, 0.0)
        return (0, str(value).lower())

    if isinstance(value, dict):
        return (0, str(value).lower()) if not is_numeric else (1, 0.0)

    if field.get("type") == "master":
        lookup = master_lookup_by_field.get(flat_key, {})
        display_key = column.get("display_key")
        if display_key:
            text = render_master_display_text(value, lookup, str(display_key))
        else:
            text = render_master_display_text(value, lookup)
        return (0, text.lower()) if text else (1, "")

    if is_numeric:
        try:
            return (0, float(value))
        except (ValueError, TypeError):
            return (1, 0.0)

    return (0, str(value).lower())


def sort_submissions(
    submissions: list[dict[str, Any]],
    sort: str,
    order: str,
    display_columns: list[dict[str, Any]],
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]],
) -> None:
    """Sort the submission list in place."""
    if order not in ("asc", "desc"):
        order = "desc"
    reverse = order == "desc"

    if sort in ("created_at", "updated_at"):
        submissions.sort(
            key=lambda s: str(s.get(sort) or ""), reverse=reverse
        )
        return

    if sort == "username":
        submissions.sort(
            key=lambda s: (
                s.get("_display_username") or s.get("username") or ""
            ).lower(),
            reverse=reverse,
        )
        return

    if sort.isdigit():
        col_idx = int(sort)
        if 0 <= col_idx < len(display_columns):
            column = display_columns[col_idx]
            submissions.sort(
                key=lambda item: _get_column_sort_key(
                    item, column, master_lookup_by_field
                ),
                reverse=reverse,
            )
            return

    # Fallback: created_at desc
    submissions.sort(key=lambda s: str(s.get("created_at") or ""), reverse=True)


def build_submission_row_values(
    data: dict[str, Any],
    display_columns: list[dict[str, Any]],
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]],
    file_names: dict[str, str],
    temporal_style: str = "raw",
) -> list[str]:
    row_values: list[str] = []
    for column in display_columns:
        field = column["field"]
        flat_key = field["flat_key"]
        value = get_nested_value(data, flat_key)

        if field.get("type") == "group" and field.get("is_array"):
            row_values.append(
                format_array_group_value(value, field.get("children", []))
            )
            continue

        if field.get("type") == "master":
            lookup = master_lookup_by_field.get(flat_key, {})
            if column["kind"] == "master_display":
                row_values.append(
                    render_master_display_text(
                        value, lookup, str(column.get("display_key", ""))
                    )
                )
            else:
                row_values.append(render_master_display_text(value, lookup))
            continue

        field_type = field.get("type", "")
        if temporal_style != "raw" and field_type in _TEMPORAL_KINDS:
            row_values.append(
                _format_temporal_value(field_type, value, iso=temporal_style == "iso")
            )
            continue

        row_values.append(value_to_text(value, file_names, field_type == "file"))
    return row_values


def _resolve_master_display_raw(
    raw_value: Any,
    lookup: dict[str, dict[str, Any]],
    display_key: str,
) -> Any:
    """参照フィールドの表示用ファイル値を生ID (またはID配列) で返す。"""

    def resolve_one(value: Any) -> Any:
        if value in (None, ""):
            return None
        row = lookup.get(str(value))
        if not row:
            return None
        return (row.get("values") or {}).get(display_key)

    if isinstance(raw_value, list):
        items: list[Any] = []
        for item in raw_value:
            resolved = resolve_one(item)
            if isinstance(resolved, list):
                items.extend(resolved)
            elif resolved not in (None, ""):
                items.append(resolved)
        return items
    return resolve_one(raw_value)


def build_submission_raw_values(
    data: dict[str, Any],
    display_columns: list[dict[str, Any]],
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]] | None = None,
) -> list[Any]:
    """ファイルフィールドのプレビュー表示などで元の値が必要な箇所向けに、
    列ごとの生の値（ID・配列など）を返す。"""
    master_lookup_by_field = master_lookup_by_field or {}
    result: list[Any] = []
    for column in display_columns:
        field = column["field"]
        flat_key = field["flat_key"]
        raw = get_nested_value(data, flat_key)
        if (
            column.get("kind") == "master_display"
            and column.get("display_type") == "file"
        ):
            lookup = master_lookup_by_field.get(flat_key, {})
            raw = _resolve_master_display_raw(
                raw, lookup, str(column.get("display_key", ""))
            )
        result.append(raw)
    return result


def collect_submission_master_display_file_ids(
    submissions: list[dict[str, Any]],
    display_columns: list[dict[str, Any]],
    master_lookup_by_field: dict[str, dict[str, dict[str, Any]]],
) -> set[str]:
    """送信一覧で参照フィールド越しに表示されるファイルの ID を収集する。"""
    ids: set[str] = set()
    file_columns = [
        column
        for column in display_columns
        if column.get("kind") == "master_display"
        and column.get("display_type") == "file"
    ]
    if not file_columns:
        return ids
    for submission in submissions:
        data = submission.get("data_json", {})
        for column in file_columns:
            flat_key = column["field"]["flat_key"]
            lookup = master_lookup_by_field.get(flat_key, {})
            raw = _resolve_master_display_raw(
                get_nested_value(data, flat_key),
                lookup,
                str(column.get("display_key", "")),
            )
            if isinstance(raw, list):
                for item in raw:
                    if isinstance(item, str) and item:
                        ids.add(item)
            elif isinstance(raw, str) and raw:
                ids.add(raw)
    return ids


async def gather_filtered_submissions(
    request: Request,
    form_id: str,
    *,
    filter_user_id: int | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, str]]:
    """フォーム・フィールド定義と、現在のクエリでフィルター済みの送信一覧、
    ファイル名マップを返す。

    送信一覧／エクスポート／集計で共通のデータ取得手順（list_submissions →
    配列グループ行の展開 → apply_filters）をまとめたもの。
    """
    storage = request.app.state.storage
    form = storage.forms.get_form(form_id)
    if not form:
        raise HTTPException(status_code=404, detail="フォームが見つかりません")

    fields = fields_from_schema(form["schema_json"], form.get("field_order", []))
    submissions = storage.submissions.list_submissions(form_id)
    if filter_user_id is not None:
        submissions = [s for s in submissions if s.get("user_id") == filter_user_id]
    expanded_submissions: list[dict[str, Any]] = []
    for submission in submissions:
        data = submission.get("data_json", {})
        for expanded_data in expand_group_array_rows(fields, data):
            expanded_submissions.append({**submission, "data_json": expanded_data})
    file_ids = collect_file_ids(submissions, fields)
    file_names = resolve_file_names(storage.files, file_ids)
    filtered = apply_filters(
        expanded_submissions, fields, dict(request.query_params), file_names=file_names
    )
    return form, fields, filtered, file_names


async def build_full_table_context(
    request: Request,
    fields: list[dict[str, Any]],
    submissions: list[dict[str, Any]],
) -> dict[str, Any]:
    """ページングなしで送信一覧と同じ表示行を構築する（集計ページの元データ表示用）。

    送信一覧の表示列／値整形ロジックをそのまま再利用し、与えられた送信集合を
    すべて行に変換する。送信ユーザーは認証プロバイダの表示名へ解決する。
    """
    storage = request.app.state.storage
    display_columns, master_lookup_by_field = build_submission_display_columns(
        storage, fields
    )
    file_ids = collect_file_ids(submissions, fields) | (
        collect_submission_master_display_file_ids(
            submissions, display_columns, master_lookup_by_field
        )
    )
    file_infos = resolve_file_infos(storage.files, file_ids, file_url_builder(request))
    file_names = {fid: info["name"] for fid, info in file_infos.items()}
    user_display_map = await resolve_user_display_map(request)

    rows: list[dict[str, Any]] = []
    for item in submissions:
        data = item.get("data_json", {})
        rows.append(
            {
                "id": item["id"],
                "created_at": item.get("created_at"),
                "updated_at": item.get("updated_at"),
                "username": resolve_user_label(item, user_display_map),
                "values": build_submission_row_values(
                    data,
                    display_columns,
                    master_lookup_by_field,
                    file_names,
                    temporal_style="display",
                ),
                "raw_values": build_submission_raw_values(
                    data, display_columns, master_lookup_by_field
                ),
            }
        )
    return {
        "display_columns": display_columns,
        "file_infos": file_infos,
        "rows": rows,
    }


async def build_submission_list_context(
    request: Request,
    form_id: str,
    *,
    include_user_display_map: bool,
    filter_user_id: int | None = None,
) -> dict[str, Any]:
    """共通の送信一覧コンテキストを構築する。"""
    storage = request.app.state.storage
    form = storage.forms.get_form(form_id)
    if not form:
        raise HTTPException(status_code=404, detail="フォームが見つかりません")

    fields = fields_from_schema(form["schema_json"], form.get("field_order", []))
    submissions = storage.submissions.list_submissions(form_id)
    if filter_user_id is not None:
        submissions = [s for s in submissions if s.get("user_id") == filter_user_id]
    expanded_submissions: list[dict[str, Any]] = []
    for submission in submissions:
        data = submission.get("data_json", {})
        for expanded_data in expand_group_array_rows(fields, data):
            expanded_submissions.append({**submission, "data_json": expanded_data})
    file_ids = collect_file_ids(submissions, fields)

    display_columns, master_lookup_by_field = build_submission_display_columns(
        storage, fields
    )
    file_ids |= collect_submission_master_display_file_ids(
        submissions, display_columns, master_lookup_by_field
    )
    file_infos = resolve_file_infos(storage.files, file_ids, file_url_builder(request))
    file_names = {fid: info["name"] for fid, info in file_infos.items()}

    filtered = apply_filters(
        expanded_submissions, fields, dict(request.query_params), file_names=file_names
    )

    if include_user_display_map:
        user_display_map = await resolve_user_display_map(request)
        for item in filtered:
            item["_display_username"] = resolve_user_label(item, user_display_map)

    sort = request.query_params.get("sort", "created_at")
    order = request.query_params.get("order", "desc")
    sort_submissions(filtered, sort, order, display_columns, master_lookup_by_field)

    try:
        page = int(request.query_params.get("page", 1))
    except (ValueError, TypeError):
        page = 1
    try:
        page_size = int(request.query_params.get("page_size", 50))
    except (ValueError, TypeError):
        page_size = 50
    page = max(1, page)
    page_size = max(1, min(page_size, 100))
    total = len(filtered)
    start = (page - 1) * page_size
    end = start + page_size
    page_items = filtered[start:end]

    filter_fields = flatten_filter_fields(fields)

    display_rows = []
    for item in page_items:
        data = item.get("data_json", {})
        row_values = build_submission_row_values(
            data,
            display_columns,
            master_lookup_by_field,
            file_names,
            temporal_style="display",
        )
        raw_values = build_submission_raw_values(
            data, display_columns, master_lookup_by_field
        )
        row = {
            "id": item["id"],
            "created_at": item["created_at"],
            "updated_at": item.get("updated_at"),
            "user_id": item.get("user_id"),
            "username": (
                item.get("_display_username")
                if include_user_display_map
                else item.get("username")
            )
            or "",
            "values": row_values,
            "raw_values": raw_values,
        }
        display_rows.append(row)

    total_pages = max(1, (total + page_size - 1) // page_size)

    return {
        "form": form,
        "fields": fields,
        "display_columns": display_columns,
        "filter_fields": filter_fields,
        "rows": display_rows,
        "file_infos": file_infos,
        "page": page,
        "page_size": page_size,
        "total": total,
        "total_pages": total_pages,
        "query": dict(request.query_params),
        "sort": sort,
        "order": order,
    }


async def perform_update_submission(
    request: Request, form_id: str, submission_id: str
) -> HTMLResponse | RedirectResponse:
    from jsonschema import Draft7Validator

    from schemaform.routes.public import save_upload
    from schemaform.utils import now_utc

    storage = request.app.state.storage
    templates = request.app.state.templates
    form = storage.forms.get_form(form_id)
    if not form:
        raise HTTPException(status_code=404, detail="フォームが見つかりません")
    existing = storage.submissions.get_submission(submission_id)
    if not existing or existing["form_id"] != form_id:
        raise HTTPException(status_code=404, detail="送信データが見つかりません")

    form_data = await request.form()
    fields = fields_from_schema(form["schema_json"], form.get("field_order", []))
    enrich_master_options(storage, fields)
    submission: dict[str, Any] = {}
    old_data = existing.get("data_json", {})

    async def collect_fields(
        field_list: list[dict[str, Any]],
        target: dict[str, Any],
        prefix: str,
        old_target: dict[str, Any],
    ) -> None:
        for field in field_list:
            key = field["key"]
            form_key = f"{prefix}{key}" if prefix else key
            field_type = field["type"]
            is_array = field.get("is_array", False)

            if field_type == "calculated":
                raw_value = form_data.get(form_key)
                target[key] = normalize_number(raw_value, False)
                continue

            if field_type == "group":
                children = field.get("children") or []
                if is_array:
                    indices: set[int] = set()
                    form_prefix = f"{form_key}."
                    for k in form_data:
                        if k.startswith(form_prefix):
                            rest = k[len(form_prefix) :]
                            parts = rest.split(".", 1)
                            if parts[0].isdigit():
                                indices.add(int(parts[0]))
                    items: list[dict[str, Any]] = []
                    old_items = old_target.get(key, []) if isinstance(old_target.get(key), list) else []
                    for order, idx in enumerate(sorted(indices)):
                        item: dict[str, Any] = {}
                        old_item = old_items[order] if order < len(old_items) and isinstance(old_items, list) else {}
                        await collect_fields(children, item, f"{form_key}.{idx}.", old_item)
                        if item:
                            items.append(item)
                    target[key] = items
                else:
                    group_data: dict[str, Any] = {}
                    old_group = old_target.get(key, {}) if isinstance(old_target.get(key), dict) else {}
                    await collect_fields(children, group_data, f"{form_key}.", old_group)
                    target[key] = group_data
                continue

            if is_array:
                if field_type == "file":
                    uploads = form_data.getlist(form_key)
                    file_ids: list[str] = []
                    has_new_upload = any(
                        upload and getattr(upload, "filename", "")
                        for upload in uploads
                    )
                    if has_new_upload:
                        for upload in uploads:
                            if upload and getattr(upload, "filename", ""):
                                file_ids.append(
                                    await save_upload(
                                        upload,
                                        form["id"],
                                        request,
                                        str(field.get("format", "")),
                                        field.get("allowed_extensions") or [],
                                    )
                                )
                        target[key] = file_ids
                    else:
                        target[key] = old_target.get(key, [])
                    continue

                values = [v for v in form_data.getlist(form_key) if v not in (None, "")]
                if field_type in {"number", "integer"}:
                    parsed = [
                        normalize_number(v, field_type == "integer")
                        for v in values
                        if normalize_number(v, field_type == "integer") is not None
                    ]
                    target[key] = parsed
                elif field_type == "boolean":
                    target[key] = [parse_bool(v) for v in values]
                else:
                    target[key] = values
            else:
                if field_type == "file":
                    upload = form_data.get(form_key)
                    if upload and getattr(upload, "filename", ""):
                        target[key] = await save_upload(
                            upload,
                            form["id"],
                            request,
                            str(field.get("format", "")),
                            field.get("allowed_extensions") or [],
                        )
                    else:
                        target[key] = old_target.get(key)
                    continue

                raw_value = form_data.get(form_key)
                if field_type in {"number", "integer"}:
                    target[key] = normalize_number(raw_value, field_type == "integer")
                elif field_type == "boolean":
                    target[key] = parse_bool(raw_value)
                else:
                    target[key] = str(raw_value) if raw_value is not None else None

    await collect_fields(fields, submission, "", old_data)
    submission = clean_empty_recursive(submission) or {}

    def _compute_calculated_edit(
        field_list: list[dict[str, Any]], data: dict[str, Any],
    ) -> None:
        for field in field_list:
            if field["type"] == "calculated" and field.get("formula"):
                result = evaluate_formula(field["formula"], data)
                if result is not None:
                    data[field["key"]] = result
            elif field["type"] == "group" and not field.get("is_array"):
                children = field.get("children") or []
                group_data = data.get(field["key"])
                if isinstance(group_data, dict):
                    _compute_calculated_edit(children, group_data)

    _compute_calculated_edit(fields, submission)

    validator = Draft7Validator(form["schema_json"])
    errors = sorted(validator.iter_errors(submission), key=lambda err: list(err.path))
    master_errors = validate_master_references(storage, fields, submission)
    if errors or master_errors:
        messages = [f"{error.message}" for error in errors] + master_errors
        file_ids = collect_file_ids([{**existing, "data_json": submission}], fields)
        file_infos = resolve_file_infos(
            storage.files, file_ids, file_url_builder(request)
        )
        return templates.TemplateResponse(
            "submission_edit.html",
            {
                "request": request,
                "form": form,
                "fields": fields,
                "submission": {**existing, "data_json": submission},
                "file_infos": file_infos,
                "errors": messages,
            },
        )

    now = now_utc()
    try:
        updated = storage.submissions.update_submission(
            submission_id, {"data_json": submission, "updated_at": now}
        )
    except KeyError:
        raise HTTPException(status_code=404, detail="送信データが見つかりません")

    if (
        form.get("webhook_url")
        and form.get("webhook_on_edit")
    ):
        from schemaform.webhook import send_webhook

        await send_webhook(form["webhook_url"], "edit", form, updated)

    return RedirectResponse(f"/forms/{form_id}/submissions", status_code=303)


_EXPORT_FORMATS = {"csv", "xlsx", "parquet", "json"}

_TEMPORAL_KINDS = {"datetime", "date", "time"}
_TEMPORAL_NUMBER_FORMATS = {
    "datetime": "yyyy/mm/dd hh:mm",
    "date": "yyyy/mm/dd",
    "time": "hh:mm",
}
_TEMPORAL_DISPLAY_FORMATS = {
    "datetime": "%Y/%m/%d %H:%M",
    "date": "%Y/%m/%d",
    "time": "%H:%M",
}
_NUMERIC_KINDS = {"number", "integer"}


def _parse_temporal(
    kind: str, text: str, *, keep_tz: bool = False
) -> datetime | date | time | None:
    """Parse a formatted cell string back into a temporal object.

    Returns None when the value is empty or not parseable so the caller can
    fall back to writing it as plain text. By default a timezone-aware datetime
    is converted to the local naive wall clock (Excel has no timezone concept);
    pass ``keep_tz=True`` to preserve the original offset for round-tripping.
    """
    text = (text or "").strip()
    if not text:
        return None
    try:
        if kind == "date":
            return date.fromisoformat(text)
        if kind == "time":
            return time.fromisoformat(text)
        if kind == "datetime":
            value = datetime.fromisoformat(text)
            if value.tzinfo is not None and not keep_tz:
                value = value.astimezone().replace(tzinfo=None)
            return value
    except ValueError:
        return None
    return None


def _format_temporal_value(kind: str, value: Any, *, iso: bool = False) -> str:
    """Normalize a stored temporal field value to a consistent notation.

    Stored values use ISO notation (e.g. ``2026-05-21T14:30``). For the
    submission list display we mirror the timestamp columns and emit slash
    notation (``2026/05/21 14:30``); for downloads we emit canonical ISO 8601
    at minute precision so every export format matches. Timezone offsets on the
    input are preserved (kept in ISO output, and the wall clock is shown as-is
    rather than shifted) so the value's meaning never changes silently.
    Unparseable values fall back to their original text.
    """
    if isinstance(value, list):
        return ", ".join(
            _format_temporal_value(kind, item, iso=iso)
            for item in value
            if item is not None
        )
    if value in (None, ""):
        return ""
    parsed = _parse_temporal(kind, str(value), keep_tz=True)
    if parsed is None:
        return str(value)
    if iso:
        if kind == "date":
            return parsed.isoformat()
        return parsed.isoformat(timespec="minutes")
    return parsed.strftime(_TEMPORAL_DISPLAY_FORMATS[kind])


def _display_text_width(text: str) -> int:
    """Approximate the rendered width of ``text`` in Excel column units.

    Full-width / wide (mostly CJK) characters take roughly two units, others
    take one, so Japanese labels are not under-sized.
    """
    width = 0
    for char in text:
        width += 2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
    return width


def _parse_number(kind: str, text: str) -> int | float | None:
    """Parse a cell string back into a number for Excel.

    Returns None when the value is empty or not numeric so the caller can fall
    back to writing it as plain text. An "integer" column may still hold a
    non-integral value, so it falls back to float rather than dropping it.
    """
    text = (text or "").strip()
    if not text:
        return None
    try:
        if kind == "integer":
            return int(text)
        return float(text)
    except ValueError:
        try:
            return float(text)
        except ValueError:
            return None


def _column_export_kind(column: dict[str, Any]) -> str:
    """Return the export kind for a display column, else "".

    Temporal kinds ("datetime"/"date"/"time") let the Excel export emit real
    date values, and numeric kinds ("number"/"integer") let it emit real
    numbers. Array columns join multiple values into one string, so they stay
    text.
    """
    field = column.get("field", {})
    if field.get("is_array"):
        return ""
    typed_kinds = _TEMPORAL_KINDS | _NUMERIC_KINDS
    if column.get("kind") == "master_display":
        display_type = column.get("display_type", "")
        return display_type if display_type in typed_kinds else ""
    field_type = field.get("type", "")
    return field_type if field_type in typed_kinds else ""


def _unique_column_names(headers: list[str]) -> list[str]:
    """Disambiguate duplicate header labels for formats that need unique keys."""
    seen: dict[str, int] = {}
    result: list[str] = []
    for header in headers:
        name = header or "column"
        if name in seen:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            result.append(name)
    return result


def _serialize_export(
    fmt: str,
    headers: list[str],
    rows: list[list[str]],
    column_kinds: list[str] | None = None,
    column_wraps: list[bool] | None = None,
) -> tuple[bytes | str, str, str]:
    """Serialize tabular data to the requested format.

    ``column_kinds`` aligns with ``headers`` and marks temporal columns
    ("datetime"/"date"/"time") so the Excel export can emit real date values.
    ``column_wraps`` aligns with ``headers`` and marks columns whose Excel
    cells should enable wrap text (multi-line text fields).

    Returns (content, media_type, file_extension).
    """
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        kinds = column_kinds or []
        wraps = column_wraps or []
        # Excel's standard vertical alignment is centered; an unset (None)
        # alignment renders at the bottom, so pin every cell to center.
        center_align = Alignment(vertical="center")
        center_wrap_align = Alignment(vertical="center", wrap_text=True)
        header_font = Font(bold=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "submissions"
        worksheet.append(headers)
        # Force header cells to text so labels like "=total" are not treated
        # as Excel formulas.
        for cell in worksheet[1]:
            cell.alignment = center_align
            cell.font = header_font
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for row in rows:
            worksheet.append(row)
            for index, cell in enumerate(worksheet[worksheet.max_row]):
                wrap = index < len(wraps) and wraps[index]
                cell.alignment = center_wrap_align if wrap else center_align
                if not isinstance(cell.value, str):
                    continue
                kind = kinds[index] if index < len(kinds) else ""
                if kind in _TEMPORAL_KINDS:
                    parsed = _parse_temporal(kind, cell.value)
                    if parsed is not None:
                        cell.value = parsed
                        cell.number_format = _TEMPORAL_NUMBER_FORMATS[kind]
                        continue
                elif kind in _NUMERIC_KINDS:
                    number = _parse_number(kind, cell.value)
                    if number is not None:
                        cell.value = number
                        continue
                # Normalize CRLF/CR to LF so Excel renders a single in-cell
                # line break instead of doubling it.
                cell.value = cell.value.replace("\r\n", "\n").replace("\r", "\n")
                # Force text so values like "=cmd" are stored as literals
                # rather than being interpreted as Excel formulas.
                cell.data_type = "s"
        # Size each column to its content. Full-width (CJK) characters count
        # as two units, and wrapped cells are measured by their longest line
        # so multi-line text does not blow up the width.
        for col_index in range(len(headers)):
            content_width = _display_text_width(headers[col_index])
            for row in rows:
                if col_index >= len(row):
                    continue
                text = row[col_index].replace("\r\n", "\n").replace("\r", "\n")
                line_width = max(
                    (_display_text_width(line) for line in text.split("\n")),
                    default=0,
                )
                content_width = max(content_width, line_width)
            # Pad for cell margins plus the header's filter dropdown arrow,
            # then clamp so columns stay within a sensible range.
            width = max(8, min(content_width + 4, 60))
            worksheet.column_dimensions[get_column_letter(col_index + 1)].width = width
        # Enable Excel's filter dropdowns on the header row over all data.
        worksheet.auto_filter.ref = worksheet.dimensions
        buffer = io.BytesIO()
        workbook.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )

    if fmt == "parquet":
        import pyarrow as pa
        import pyarrow.parquet as pq

        columns = _unique_column_names(headers)
        arrays = [
            pa.array([row[index] for row in rows], type=pa.string())
            for index in range(len(columns))
        ]
        table = pa.Table.from_arrays(arrays, names=columns)
        buffer = pa.BufferOutputStream()
        pq.write_table(table, buffer)
        return buffer.getvalue().to_pybytes(), "application/vnd.apache.parquet", "parquet"

    if fmt == "json":
        columns = _unique_column_names(headers)
        records = [dict(zip(columns, row)) for row in rows]
        return (
            json.dumps(records, ensure_ascii=False, indent=2),
            "application/json; charset=utf-8",
            "json",
        )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue(), "text/csv; charset=utf-8", "csv"


def _parse_correct_map(raw: str | None) -> dict[str, dict[str, Any]]:
    """集計ページから渡される正解指定(JSON)をパースする。

    形式: {flat_key: {"mode": "single", "value": str}}
          {flat_key: {"mode": "array", "list": [str], "ordered": bool}}
    enumフィールドの正答数・正答率をダウンロードに含めるために使う。"""
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, ValueError, TypeError):
        return {}
    if not isinstance(parsed, dict):
        return {}
    result: dict[str, dict[str, Any]] = {}
    for key, info in parsed.items():
        if not isinstance(info, dict):
            continue
        if info.get("mode") == "array":
            values = info.get("list")
            if isinstance(values, list) and values:
                result[str(key)] = {
                    "mode": "array",
                    "list": [str(v) for v in values],
                    "ordered": bool(info.get("ordered")),
                }
        else:
            value = info.get("value")
            if value not in (None, ""):
                result[str(key)] = {"mode": "single", "value": str(value)}
    return result


def _row_is_correct(
    data: dict[str, Any], flat_key: str, info: dict[str, Any]
) -> bool:
    value = get_nested_value(data, flat_key)
    if value is None:
        return False
    if info["mode"] == "array":
        actual = [str(v) for v in (value if isinstance(value, list) else [value])]
        expected = info["list"]
        if len(actual) != len(expected):
            return False
        if info.get("ordered"):
            return actual == expected
        return sorted(actual) == sorted(expected)
    if isinstance(value, list):
        return False
    return str(value) == info["value"]


def _correctness_cells(
    data: dict[str, Any], correct_map: dict[str, dict[str, Any]]
) -> tuple[str, str]:
    """1送信あたりの (正答数, 正答率) セル文字列を返す。画面表示と同じ表記。"""
    total = len(correct_map)
    if total == 0:
        return "", ""
    correct = sum(
        1 for key, info in correct_map.items() if _row_is_correct(data, key, info)
    )
    return f"{correct} / {total}", f"{round(correct / total * 100)}%"


@router.get("/forms/{form_id}/export", tags=["admin"])
async def export_submissions(
    request: Request, form_id: str, _: Any = Depends(form_editor_guard)
) -> Response:
    storage = request.app.state.storage
    form, fields, filtered, file_names = await gather_filtered_submissions(
        request, form_id
    )
    display_columns, master_lookup_by_field = build_submission_display_columns(
        storage, fields
    )

    user_display_map = await resolve_user_display_map(request)
    for item in filtered:
        item["_display_username"] = resolve_user_label(item, user_display_map)

    sort = request.query_params.get("sort", "created_at")
    order = request.query_params.get("order", "desc")
    sort_submissions(filtered, sort, order, display_columns, master_lookup_by_field)

    correct_map = _parse_correct_map(request.query_params.get("correct"))
    include_correct = bool(correct_map)

    def _fmt(value: Any) -> str:
        if isinstance(value, datetime):
            return value.astimezone().strftime("%Y-%m-%dT%H:%M")
        return str(value or "")

    correct_headers = ["正答数", "正答率"] if include_correct else []
    headers = (
        ["送信日時", "更新日時", "送信ユーザー"]
        + correct_headers
        + [column["label"] for column in display_columns]
    )
    column_kinds = (
        ["datetime", "datetime", ""]
        + ([""] * len(correct_headers))
        + [_column_export_kind(column) for column in display_columns]
    )
    column_wraps = (
        [False, False, False]
        + ([False] * len(correct_headers))
        + [bool(column.get("field", {}).get("multiline")) for column in display_columns]
    )
    rows = []
    for submission in filtered:
        data = submission.get("data_json", {})
        row = [
            _fmt(submission.get("created_at")),
            _fmt(submission.get("updated_at")),
            submission.get("_display_username") or "",
        ]
        if include_correct:
            row += list(_correctness_cells(data, correct_map))
        row += build_submission_row_values(
            data,
            display_columns,
            master_lookup_by_field,
            file_names,
            temporal_style="iso",
        )
        rows.append(row)

    fmt = request.query_params.get("format", "csv")
    if fmt not in _EXPORT_FORMATS:
        fmt = "csv"

    content, content_type, extension = _serialize_export(
        fmt, headers, rows, column_kinds, column_wraps
    )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    raw_name = (form.get("name") or "submissions").strip() or "submissions"
    safe_name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", raw_name)
    filename = f"{safe_name}_{timestamp}.{extension}"
    ascii_fallback = (
        re.sub(r"[^A-Za-z0-9._-]+", "_", filename) or f"submissions.{extension}"
    )
    return Response(
        content=content,
        media_type=content_type,
        headers={
            "Content-Disposition": (
                f"attachment; filename=\"{ascii_fallback}\"; "
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )


def _build_import_field_map(
    fields: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    """Build a mapping from label/key to flat field info for CSV import."""
    flat_fields = flatten_fields(fields, expand_rows_for_group_arrays=True)
    mapping: dict[str, dict[str, Any]] = {}
    for field in flat_fields:
        mapping[field["flat_label"]] = field
        mapping[field["flat_key"]] = field
    return mapping


def _wrap_arrays_from_schema(
    fields: list[dict[str, Any]], data: dict[str, Any]
) -> None:
    """Wrap dict values into single-element lists for array group fields."""
    for field in fields:
        key = field["key"]
        if key not in data:
            continue
        if field.get("type") == "group":
            if field.get("is_array"):
                if isinstance(data[key], dict):
                    data[key] = [data[key]]
            else:
                children = field.get("children") or []
                if isinstance(data[key], dict) and children:
                    _wrap_arrays_from_schema(children, data[key])


def _convert_cell_value(raw: str, field: dict[str, Any]) -> Any:
    """Convert a raw CSV cell string to the appropriate Python type."""
    field_type = field.get("type", "string")
    if raw == "":
        return None
    if field.get("is_array") or field_type == "group":
        try:
            parsed = json.loads(raw)
            if isinstance(parsed, list):
                return parsed
        except (json.JSONDecodeError, ValueError):
            pass
    if field_type in ("number", "integer"):
        return normalize_number(raw, field_type == "integer")
    if field_type == "boolean":
        return parse_bool(raw)
    return raw


@router.post(
    "/forms/{form_id}/import", tags=["admin"]
)
async def import_submissions(
    request: Request, form_id: str, _: Any = Depends(form_editor_guard)
) -> RedirectResponse:
    from jsonschema import Draft7Validator

    storage = request.app.state.storage
    form = storage.forms.get_form(form_id)
    if not form:
        raise HTTPException(status_code=404, detail="フォームが見つかりません")

    form_data = await request.form()
    upload = form_data.get("file")
    if not upload or not getattr(upload, "filename", ""):
        raise HTTPException(status_code=400, detail="ファイルを選択してください")

    content = await upload.read()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("shift_jis")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=400, detail="ファイルのエンコーディングを認識できません"
            )

    filename = getattr(upload, "filename", "") or ""
    if filename.lower().endswith(".tsv"):
        delimiter = "\t"
    else:
        delimiter = ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        headers = next(reader)
    except StopIteration:
        raise HTTPException(status_code=400, detail="ファイルが空です")

    fields = fields_from_schema(form["schema_json"], form.get("field_order", []))
    field_map = _build_import_field_map(fields)

    col_field_map: list[dict[str, Any] | None] = []
    for header in headers:
        header_stripped = header.strip()
        col_field_map.append(field_map.get(header_stripped))

    validator = Draft7Validator(form["schema_json"])

    now = now_utc()
    imported_count = 0
    skipped_count = 0
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        data: dict[str, Any] = {}
        for col_idx, cell in enumerate(row):
            if col_idx >= len(col_field_map):
                break
            field_info = col_field_map[col_idx]
            if field_info is None:
                continue
            flat_key = field_info["flat_key"]
            value = _convert_cell_value(cell.strip(), field_info)
            if value is not None:
                set_nested_value(data, flat_key, value)

        _wrap_arrays_from_schema(fields, data)
        data = clean_empty_recursive(data) or {}
        if not data:
            continue

        schema_errors = list(validator.iter_errors(data))
        master_errors = validate_master_references(storage, fields, data)
        if schema_errors or master_errors:
            skipped_count += 1
            continue

        submission_id = new_ulid()
        storage.submissions.create_submission(
            {
                "id": submission_id,
                "form_id": form_id,
                "data_json": data,
                "created_at": now,
            }
        )
        imported_count += 1

    return RedirectResponse(f"/forms/{form_id}/submissions", status_code=303)


@router.get("/healthz", tags=["system"])
async def healthz() -> dict[str, str]:
    return {"status": "ok"}
